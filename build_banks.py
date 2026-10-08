#!/usr/bin/env python3
"""Build the word banks for Key Terms.

Teachers edit banks/terms.xlsx (one tab per subject). Terms sent in
through the teacher site arrive as CSV files in banks/inbox/<subject-id>/
and are read after that subject's tab, exactly as if they were extra rows.

Run:  python build_banks.py          check, then write banks/banks.js
      python build_banks.py --check  check only, write nothing
(needs openpyxl: pip install openpyxl)
Writing also updates the cache list and version in sw.js so installed
copies pick up changes.
"""
import csv, hashlib, json, re, sys, unicodedata
from pathlib import Path

ROOT = Path(__file__).parent
XLSX = ROOT / "banks" / "terms.xlsx"
INBOX = ROOT / "banks" / "inbox"
OUT = ROOT / "banks" / "banks.js"
SW = ROOT / "sw.js"
COLS = ["term", "clue", "image", "topic", "tier", "question",
        "correct_answer", "wrong_1", "wrong_2", "wrong_3", "explanation", "checked_by"]

errors, warnings = [], []

def cell_text(v):
    if v is None: return ""
    if isinstance(v, float) and v.is_integer(): v = int(v)   # 23.0 typed as a number -> "23"
    return str(v).strip()

def read_sheet(ws):
    rows = list(ws.iter_rows(values_only=True))
    if not rows: return []
    head = [cell_text(h).lower() for h in rows[0]]
    return [{head[i]: cell_text(v) for i, v in enumerate(r) if i < len(head) and head[i]} for r in rows[1:]]

def follow_up(where, r):
    opts = [r["correct_answer"], r["wrong_1"], r["wrong_2"], r["wrong_3"]]
    missing = [k for k in ("question", "correct_answer", "wrong_1", "wrong_2", "wrong_3", "explanation") if not r[k]]
    if missing:
        errors.append(f"{where}: missing {', '.join(missing)}"); return None
    if len(set(opts)) < 4:  # exact match, so capital-letter distractors are allowed
        errors.append(f"{where}: the four answer options must all be different")
    return {"q": r["question"], "options": opts, "answer": 0, "explain": r["explanation"]}

def read_inbox(sid):
    """Rows from banks/inbox/<sid>/*.csv, oldest file first, as (where, row)."""
    out = []
    for f in sorted((INBOX / sid).glob("*.csv")):
        with f.open(newline="", encoding="utf-8-sig") as fh:
            for n, row in enumerate(csv.DictReader(fh), start=2):
                row = {(k or "").strip().lower(): (v or "").strip() for k, v in row.items()}
                out.append((f"inbox file {f.relative_to(ROOT).as_posix()}, row {n}", row))
    return out

def check_row(where, row, by_key, items):
    """A term's first row defines it. Later rows with the same term and a
    blank clue add extra follow-up questions, one of which is picked at random."""
    r = {k: (row.get(k) or "").strip() for k in COLS}
    term = r["term"]
    if not term:
        errors.append(f"{where}: missing term"); return
    plain = unicodedata.normalize("NFD", term).encode("ascii", "ignore").decode()
    key = re.sub(r"[^a-z]", "", plain.lower())

    if key in by_key:
        if r["clue"]:
            errors.append(f"{where}: '{term}' already appears ({by_key[key]['row']}). "
                          "To add another follow-up question, leave the clue blank.")
            return
        fu = follow_up(where, r)
        if fu: by_key[key]["item"]["followups"].append(fu)
        return

    missing = [k for k in ("clue", "topic", "tier") if not r[k]]
    if missing:
        errors.append(f"{where}: first row for '{term}' is missing {', '.join(missing)}"); return
    if re.search(r"[^A-Za-zÀ-ÖØ-öø-ÿ \-]", term):
        errors.append(f"{where}: term '{term}' may only contain letters (accents allowed), spaces and hyphens")
    if not key:
        errors.append(f"{where}: term has no letters")
    if r["tier"].upper() not in ("F", "H"):
        errors.append(f"{where}: tier must be F or H, not '{r['tier']}'")
    if r["image"]:
        if not (ROOT / r["image"]).exists():
            errors.append(f"{where}: image '{r['image']}' not found")
    elif re.search(r"\b" + re.escape(term.lower()) + r"\b", r["clue"].lower()):
        warnings.append(f"{where}: the clue contains the answer '{term}'")
    fu = follow_up(where, r)
    item = {"term": term, "clue": r["clue"], "image": r["image"] or None,
            "topic": r["topic"], "tier": r["tier"].upper(), "followups": [fu] if fu else []}
    by_key[key] = {"row": where, "item": item}
    items.append(item)

def main():
    check_only = "--check" in sys.argv[1:]
    try:
        from openpyxl import load_workbook
    except ImportError:
        sys.exit("openpyxl is needed: pip install openpyxl")
    wb = load_workbook(XLSX, data_only=True)
    if "Subjects" not in wb.sheetnames:
        sys.exit("terms.xlsx has no 'Subjects' tab")
    banks, unchecked, inbox_rows, known = [], 0, 0, set()
    for s in read_sheet(wb["Subjects"]):
        sid, name = s.get("id", ""), s.get("subject", "")
        if not sid: continue
        known.add(sid)
        if name not in wb.sheetnames:
            errors.append(f"Subjects tab lists '{name}' but there is no tab with that name"); continue
        missing = [c for c in ("term", "clue", "question", "correct_answer") if c not in
                   [cell_text(h).lower() for h in next(wb[name].iter_rows(max_row=1, values_only=True))]]
        if missing:
            errors.append(f"'{name}' tab is missing column headings: {', '.join(missing)}"); continue
        by_key, items = {}, []
        rows = [(f"'{name}' tab, row {n}", row) for n, row in enumerate(read_sheet(wb[name]), start=2)]
        extra = read_inbox(sid)
        inbox_rows += len(extra)
        for where, row in rows + extra:
            if not any(row.values()):
                continue  # blank row
            if row.get("clue") and not row.get("checked_by"): unchecked += 1
            check_row(where, row, by_key, items)
        if len(items) < 10:
            warnings.append(f"{name}: only {len(items)} terms; sets are 10 long")
        banks.append({"id": sid, "subject": name, "department": s.get("department", ""), "items": items})

    if INBOX.is_dir():
        for d in INBOX.iterdir():
            if d.is_dir() and d.name not in known and any(d.glob("*.csv")):
                errors.append(f"inbox folder '{d.name}' doesn't match any id in the Subjects tab")
    if inbox_rows:
        print(f"Included {inbox_rows} rows from banks/inbox (run merge_inbox.py to move them into terms.xlsx).")
    if unchecked:
        print(f"Note: {unchecked} terms have no checked_by initials yet.")
    for w in warnings: print("WARNING:", w)
    if errors:
        for e in errors: print("ERROR:", e)
        print(f"\n{len(errors)} error(s). Nothing was written. Fix them and run again.")
        sys.exit(1)
    if check_only:
        total = sum(len(b["items"]) for b in banks)
        print(f"Check passed: {len(banks)} subjects, {total} terms. Nothing written (--check).")
        return

    js = "// Generated by build_banks.py from banks/terms.xlsx. Don't edit by hand.\n"
    js += "window.BANKS = " + json.dumps(banks, ensure_ascii=False, indent=1) + ";\n"
    OUT.write_text(js, encoding="utf-8")

    # Refresh the service worker's file list and version.
    files = ["./", "./index.html", "./play.html", "./manifest.json", "./config.js", "./banks/banks.js"]
    files += sorted("./" + p.relative_to(ROOT).as_posix() for p in (ROOT / "icons").glob("*.png"))
    files += sorted("./" + p.relative_to(ROOT).as_posix() for p in (ROOT / "images").glob("*.svg"))
    h = hashlib.sha1()
    for f in files:
        p = ROOT / (f[2:] or "index.html")
        if p.is_file(): h.update(p.read_bytes())
    block = (f'const VERSION = "key-terms-{h.hexdigest()[:10]}";\n'
             f"const FILES = {json.dumps(files, indent=2)};\n")
    sw = SW.read_text(encoding="utf-8")
    sw = re.sub(r"// BUILD:START\n.*?// BUILD:END", "// BUILD:START\n" + block + "// BUILD:END", sw, flags=re.S)
    SW.write_text(sw, encoding="utf-8")

    total = sum(len(b["items"]) for b in banks)
    print(f"Built {len(banks)} subjects, {total} terms -> {OUT.relative_to(ROOT)}")
    for b in banks:
        fus = sum(len(i["followups"]) for i in b["items"])
        print(f"  {b['subject']}: {len(b['items'])} terms, {fus} follow-ups ({sum(1 for i in b['items'] if i['image'])} with pictures)")

if __name__ == "__main__":
    main()
