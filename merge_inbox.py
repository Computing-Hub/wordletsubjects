#!/usr/bin/env python3
"""Move terms sent in through the teacher site into banks/terms.xlsx.

Approved submissions sit in banks/inbox/<subject-id>/*.csv. The build already
includes them, so this step is only housekeeping: run it now and then on your
own computer to fold them into the workbook, where they're easier to check
and edit.

Run:  python merge_inbox.py
Then: python build_banks.py, and commit (the inbox files are deleted).

New topics are added to the Lists tab so the dropdowns keep working.
Close terms.xlsx in Excel before running this, or the save will fail.
"""
import csv, re, subprocess, sys
from pathlib import Path

ROOT = Path(__file__).parent
XLSX = ROOT / "banks" / "terms.xlsx"
INBOX = ROOT / "banks" / "inbox"
COLS = ["term", "clue", "image", "topic", "tier", "question",
        "correct_answer", "wrong_1", "wrong_2", "wrong_3", "explanation", "checked_by"]


def last_used_row(ws):
    for r in range(ws.max_row, 1, -1):
        if any(ws.cell(r, c).value not in (None, "") for c in range(1, len(COLS) + 1)):
            return r
    return 1


def grow_ranges(ws, last_row):
    """Stretch dropdowns and red-cell highlighting down to cover the new rows."""
    for dv in ws.data_validations.dataValidation:
        new = []
        for rng in str(dv.sqref).split():
            m = re.fullmatch(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", rng)
            if m and int(m.group(4)) < last_row:
                rng = f"{m.group(1)}{m.group(2)}:{m.group(3)}{last_row + 50}"
            new.append(rng)
        dv.sqref = " ".join(new)
    cf = ws.conditional_formatting
    grown = []
    for rng in list(cf):
        m = re.fullmatch(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", str(rng.sqref))
        if m and int(m.group(4)) < last_row:
            grown.append((f"{m.group(1)}{m.group(2)}:{m.group(3)}{last_row + 50}", rng.rules))
            del cf._cf_rules[rng]
    for ref, rules in grown:
        for rule in rules:
            cf.add(ref, rule)


def add_topic(wb, subject, topic):
    """Add a topic to this subject's column in the Lists tab if it's new."""
    ws = wb["Lists"]
    head = [c.value for c in ws[1]]
    if subject not in head:
        return False
    col = head.index(subject) + 1
    values = [ws.cell(r, col).value for r in range(2, ws.max_row + 2)]
    if topic in values:
        return False
    r = 2
    while ws.cell(r, col).value not in (None, ""):
        r += 1
    ws.cell(r, col, topic)
    # Make sure the topic dropdown on the subject tab reaches the new cell.
    letter = ws.cell(1, col).column_letter
    for dv in wb[subject].data_validations.dataValidation:
        m = re.fullmatch(rf"=Lists!\${letter}\$2:\${letter}\$(\d+)", dv.formula1 or "")
        if m and int(m.group(1)) < r:
            dv.formula1 = f"=Lists!${letter}$2:${letter}${r + 20}"
    return True


def main():
    files = sorted(INBOX.glob("*/*.csv")) if INBOX.is_dir() else []
    if not files:
        print("The inbox is empty. Nothing to merge.")
        return
    try:
        from openpyxl import load_workbook
    except ImportError:
        sys.exit("openpyxl is needed: pip install openpyxl")
    wb = load_workbook(XLSX)
    ids = {}
    for row in wb["Subjects"].iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            ids[str(row[0]).strip()] = str(row[1]).strip()

    merged, new_topics = 0, []
    for f in files:
        sid = f.parent.name
        if sid not in ids or ids[sid] not in wb.sheetnames:
            sys.exit(f"{f}: no subject with id '{sid}' in the Subjects tab. Nothing was changed.")
        ws = wb[ids[sid]]
        head = [str(c.value or "").strip().lower() for c in ws[1]]
        with f.open(newline="", encoding="utf-8-sig") as fh:
            rows = [{(k or "").strip().lower(): (v or "").strip() for k, v in r.items()} for r in csv.DictReader(fh)]
        r = last_used_row(ws)
        for row in rows:
            r += 1
            for key in COLS:
                if key in head:
                    ws.cell(r, head.index(key) + 1, row.get(key) or None)
            if row.get("clue") and row.get("topic") and add_topic(wb, ids[sid], row["topic"]):
                new_topics.append(f"{ids[sid]}: {row['topic']}")
            merged += 1
        grow_ranges(ws, r)

    try:
        wb.save(XLSX)
    except PermissionError:
        sys.exit("Couldn't save terms.xlsx. Close it in Excel and run this again. Nothing was changed.")
    for f in files:
        f.unlink()
    for d in INBOX.iterdir():
        if d.is_dir() and not any(d.iterdir()):
            d.rmdir()
    print(f"Moved {merged} rows from {len(files)} inbox file(s) into terms.xlsx.")
    for t in new_topics:
        print("  New topic added to the Lists tab:", t)
    print("Checking the result...")
    sys.exit(subprocess.call([sys.executable, str(ROOT / "build_banks.py"), "--check"]))


if __name__ == "__main__":
    main()
